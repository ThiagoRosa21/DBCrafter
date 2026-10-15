"""Bounded, stateless file-to-SQL conversion. Never executes generated SQL."""

import csv
import io
import json
import math
import re
import time
import unicodedata
import zipfile
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

import openpyxl
import xlrd

DIALECTS = {"sqlite", "postgresql", "mysql", "oracle"}
TYPES = {"text", "integer", "decimal", "boolean", "date", "datetime"}
IDENTIFIER = re.compile(r"[A-Za-z_][A-Za-z0-9_]{0,62}\Z", re.ASCII)
NUMBER = re.compile(r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?\Z")


class ConversionError(ValueError):
    """A safe validation failure; never includes cell contents."""


@dataclass(frozen=True)
class Limits:
    upload_bytes: int = 5 * 1024 * 1024
    output_bytes: int = 20 * 1024 * 1024
    rows: int = 10_000
    columns: int = 100
    cell_chars: int = 32_767
    cells: int = 250_000
    zip_bytes: int = 20 * 1024 * 1024
    zip_entries: int = 1_000
    seconds: float = 10


class Budget:
    def __init__(self, seconds):
        self.deadline = time.monotonic() + seconds

    def check(self):
        if time.monotonic() > self.deadline:
            raise ConversionError(
                "Tempo de conversão excedido. Divida o arquivo em partes menores."
            )


def normalize_columns(headers, limits):
    if not headers or len(headers) > limits.columns:
        raise ConversionError(f"Use de 1 a {limits.columns} colunas.")
    names = []
    for index, header in enumerate(headers, 1):
        if not isinstance(header, str) or len(header) > limits.cell_chars:
            raise ConversionError(f"O cabeçalho da coluna {index} deve ser um texto válido.")
        plain = unicodedata.normalize("NFKD", header).encode("ascii", "ignore").decode()
        name = re.sub(r"[^a-zA-Z0-9_]+", "_", plain.strip()).strip("_").lower()
        if name and name[0].isdigit():
            name = "c_" + name
        if not IDENTIFIER.fullmatch(name):
            raise ConversionError(
                f"Nome inválido na coluna {index}: use até 63 caracteres após normalização."
            )
        if name in names:
            raise ConversionError(
                "Os cabeçalhos geram nomes de colunas duplicados. Renomeie-os no arquivo."
            )
        names.append(name)
    return names


def quote_identifier(name, dialect):
    if not isinstance(name, str) or not IDENTIFIER.fullmatch(name):
        raise ConversionError(
            "Use nomes com letras ASCII, números e _, sem iniciar com número (máximo 63)."
        )
    if dialect == "oracle":
        name = name.upper()
    quote = "`" if dialect == "mysql" else '"'
    return quote + name + quote


def quote_table(table, dialect):
    parts = table.split(".")
    if not 1 <= len(parts) <= 2:
        raise ConversionError("Informe tabela ou schema.tabela.")
    return ".".join(quote_identifier(part, dialect) for part in parts)


def unique_object(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise ConversionError("O JSON contém chaves duplicadas.")
        result[key] = value
    return result


def reject_constant(_value):
    raise ConversionError("NaN e Infinity não são aceitos.")


def parse_types(raw):
    try:
        mapping = json.loads(raw or "{}", object_pairs_hook=unique_object)
    except (ValueError, RecursionError):
        raise ConversionError("O mapa de tipos deve ser um objeto JSON válido.") from None
    if not isinstance(mapping, dict) or any(
        not isinstance(value, str) or value not in TYPES for value in mapping.values()
    ):
        raise ConversionError("Tipos permitidos: text, integer, decimal, boolean, date e datetime.")
    return mapping


def validate_cell(value, limits):
    if isinstance(value, str):
        if len(value) > limits.cell_chars:
            raise ConversionError(f"Uma célula excede {limits.cell_chars} caracteres.")
        if any(
            (ord(c) < 32 and c not in "\t\r\n") or 0xD800 <= ord(c) <= 0xDFFF or ord(c) == 127
            for c in value
        ):
            raise ConversionError("Há caracteres de controle ou Unicode inválido nos dados.")
    elif isinstance(value, (dict, list)):
        raise ConversionError(
            "Use valores simples nas células; objetos e listas aninhados não são aceitos."
        )


def cast_value(value, target):
    if value is None or target is None:
        return value
    if target == "text":
        return value.isoformat() if isinstance(value, (date, datetime)) else str(value)
    if target == "date" and isinstance(value, date):
        if not isinstance(value, datetime):
            return value
        if value.time() == datetime.min.time() and value.tzinfo is None:
            return value.date()
    if target == "datetime" and isinstance(value, datetime):
        return value
    if target == "integer" and isinstance(value, (float, Decimal)):
        number = Decimal(str(value))
        if (
            number.is_finite()
            and abs(number.adjusted()) <= 100
            and number == number.to_integral_value()
        ):
            return int(number)
    raw = str(value).strip()
    try:
        if target == "integer" and len(raw) <= 100 and re.fullmatch(r"[+-]?[0-9]+", raw):
            return int(raw)
        if target == "decimal" and len(raw) <= 128 and NUMBER.fullmatch(raw):
            return Decimal(raw)
        if target == "boolean":
            if raw.lower() in {"true", "1"}:
                return True
            if raw.lower() in {"false", "0"}:
                return False
        if target == "date" and re.fullmatch(r"[0-9]{4}-[0-9]{2}-[0-9]{2}", raw):
            return date.fromisoformat(raw)
        if target == "datetime" and "T" in raw:
            return datetime.fromisoformat(raw)
    except (ValueError, InvalidOperation, OverflowError):
        pass
    raise ConversionError(
        "Um valor não corresponde ao tipo declarado; use datas ISO e ponto decimal."
    )


def text_literal(value, dialect):
    if dialect == "mysql":
        # Hex avoids dependence on NO_BACKSLASH_ESCAPES and client encodings.
        return "CONVERT(X'" + value.encode("utf-8").hex() + "' USING utf8mb4)"
    if dialect == "postgresql":
        escaped = value.replace("\\", "\\\\").replace("'", "''")
        escaped = escaped.replace("\r", "\\r").replace("\n", "\\n").replace("\t", "\\t")
        return "E'" + escaped + "'"
    if dialect == "oracle" and value == "":
        raise ConversionError(
            "Oracle trata texto vazio como NULL. Defina o tratamento no arquivo antes de exportar."
        )
    special = "\r\n\t&" if dialect == "oracle" else "\r\n\t"
    function = "CHR" if dialect == "oracle" else "char"
    pieces = []
    for part in re.split("([" + re.escape(special) + "])", value):
        if len(part) == 1 and part in special:
            pieces.append(f"{function}({ord(part)})")
        elif part:
            pieces.append("'" + part.replace("'", "''") + "'")
    return " || ".join(pieces) or "''"


def sql_literal(value, dialect):
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return ("TRUE" if value else "FALSE") if dialect == "postgresql" else str(int(value))
    if isinstance(value, (int, float, Decimal)):
        if isinstance(value, float) and not math.isfinite(value):
            raise ConversionError("Números não finitos não são aceitos.")
        if isinstance(value, Decimal) and not value.is_finite():
            raise ConversionError("Números não finitos não são aceitos.")
        if isinstance(value, Decimal) and abs(value.adjusted()) > 308:
            raise ConversionError("Expoente numérico fora dos limites de exportação.")
        literal = format(value, "f") if isinstance(value, Decimal) else str(value)
        if len(literal) > 128 or not NUMBER.fullmatch(literal):
            raise ConversionError("Número fora dos limites de exportação.")
        if abs(Decimal(literal).adjusted()) > 308:
            raise ConversionError("Expoente numérico fora dos limites de exportação.")
        return literal
    if isinstance(value, datetime):
        aware = value.utcoffset() is not None
        if aware and dialect not in {"postgresql", "sqlite"}:
            raise ConversionError(
                "Datas com fuso exigem PostgreSQL, SQLite ou conversão explícita para text."
            )
        literal = "'" + value.isoformat(sep=" ") + "'"
        if dialect == "sqlite":
            return literal
        prefix = "TIMESTAMP WITH TIME ZONE " if aware else "TIMESTAMP "
        return prefix + literal
    if isinstance(value, date):
        literal = "'" + value.isoformat() + "'"
        return literal if dialect == "sqlite" else "DATE " + literal
    if isinstance(value, str):
        return text_literal(value, dialect)
    raise ConversionError(
        "Tipo de célula não suportado; declare text para convertê-lo explicitamente."
    )


def json_table(data, budget):
    text = data.decode("utf-8-sig")
    depth = 0
    quoted = escaped = False
    for index, char in enumerate(text):
        if index % 8192 == 0:
            budget.check()
        if quoted:
            if escaped:
                escaped = False
            elif char == "\\":
                escaped = True
            elif char == '"':
                quoted = False
        elif char == '"':
            quoted = True
        elif char in "[{":
            depth += 1
            if depth > 8:
                raise ConversionError("O JSON possui aninhamento excessivo.")
        elif char in "]}":
            depth -= 1
    rows = json.loads(
        text, parse_float=Decimal, parse_constant=reject_constant, object_pairs_hook=unique_object
    )
    budget.check()
    if not isinstance(rows, list) or not rows or not isinstance(rows[0], dict) or not rows[0]:
        raise ConversionError("Use um array JSON não vazio de objetos com as mesmas chaves.")
    headers = list(rows[0])
    expected = set(headers)

    def values():
        for row in rows:
            if not isinstance(row, dict) or set(row) != expected:
                raise ConversionError("Todos os objetos JSON devem ter as mesmas chaves.")
            yield [row[key] for key in headers]

    return headers, values()


def inspect_xlsx(data, limits, budget):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if (
            len(entries) > limits.zip_entries
            or sum(item.file_size for item in entries) > limits.zip_bytes
        ):
            raise ConversionError("O Excel excede o limite de conteúdo descompactado.")
        if len({item.filename for item in entries}) != len(entries):
            raise ConversionError("O Excel contém entradas duplicadas.")
        if any(item.flag_bits & 1 for item in entries):
            raise ConversionError("Arquivos Excel criptografados não são aceitos.")
        if any(item.filename.lower().endswith("vbaproject.bin") for item in entries):
            raise ConversionError("Envie uma planilha sem macros.")
        budget.check()


@contextmanager
def read_table(data, extension, limits, budget, delimiter, null_marker):
    if extension == "csv":
        reader = csv.reader(
            io.StringIO(data.decode("utf-8-sig"), newline=""), delimiter=delimiter, strict=True
        )
        headers = next(reader, [])
        rows = ([None if cell == null_marker else cell for cell in row] for row in reader)
        yield headers, rows
    elif extension == "json":
        yield json_table(data, budget)
    elif extension == "xlsx":
        inspect_xlsx(data, limits, budget)
        book = openpyxl.load_workbook(
            io.BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
        try:
            sheet = book.worksheets[0]
            if sheet.max_column and sheet.max_column > limits.columns:
                raise ConversionError(f"Use no máximo {limits.columns} colunas.")
            sheet.reset_dimensions()

            def values():
                for row in sheet.iter_rows():
                    budget.check()
                    if len(row) > limits.columns:
                        raise ConversionError(f"Use no máximo {limits.columns} colunas.")
                    if any(cell.data_type in {"f", "e"} for cell in row):
                        raise ConversionError(
                            "Converta fórmulas e erros do Excel em valores antes do upload."
                        )
                    yield [cell.value for cell in row]

            rows = values()
            headers = next(rows, [])
            yield headers, rows
        finally:
            book.close()
    else:
        if not data.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
            raise ConversionError("O arquivo não é uma planilha XLS válida.")
        book = xlrd.open_workbook(file_contents=data, on_demand=True, logfile=io.StringIO())
        try:
            sheet = book.sheet_by_index(0)
            if sheet.ncols > limits.columns:
                raise ConversionError(f"Use no máximo {limits.columns} colunas.")

            def values():
                for index in range(sheet.nrows):
                    budget.check()
                    result = []
                    for cell in sheet.row(index):
                        if cell.ctype == xlrd.XL_CELL_ERROR:
                            raise ConversionError("Corrija as células com erro na planilha XLS.")
                        if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
                            value = None
                        elif cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(cell.value, book.datemode)
                        elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                            value = bool(cell.value)
                        else:
                            value = cell.value
                        result.append(value)
                    yield result

            rows = values()
            yield next(rows, []), rows
        finally:
            book.release_resources()


def convert(
    data,
    extension,
    table,
    dialect="sqlite",
    row_limit=None,
    types=None,
    delimiter=",",
    null_marker="\\N",
    limits=None,
):
    limits = limits or Limits()
    budget = Budget(limits.seconds)
    if not data or len(data) > limits.upload_bytes:
        raise ConversionError("Arquivo vazio ou maior que o limite de upload.")
    if extension not in {"csv", "json", "xlsx", "xls"} or dialect not in DIALECTS:
        raise ConversionError("Formato de arquivo ou banco de destino não suportado.")
    if row_limit is not None and (type(row_limit) is not int or not 1 <= row_limit <= limits.rows):
        raise ConversionError(f"O limite deve ser um inteiro entre 1 e {limits.rows}.")
    if delimiter not in {",", ";", "\t"} or not null_marker or len(null_marker) > 32:
        raise ConversionError("Separador ou marcador de NULL inválido.")
    target = quote_table(table, dialect)
    types = types or {}
    output = io.BytesIO()
    count = 0
    try:
        with read_table(data, extension, limits, budget, delimiter, null_marker) as (headers, rows):
            names = normalize_columns(headers, limits)
            if set(types) - set(names) or any(value not in TYPES for value in types.values()):
                raise ConversionError(
                    "O mapa de tipos deve usar nomes normalizados existentes e tipos permitidos."
                )
            columns = ", ".join(quote_identifier(name, dialect) for name in names)
            for row in rows:
                budget.check()
                count += 1
                if count > limits.rows or count * len(names) > limits.cells:
                    raise ConversionError(
                        "Quantidade de linhas ou células acima do limite. Divida o arquivo."
                    )
                if extension == "xlsx" and len(row) < len(names):
                    row += [None] * (len(names) - len(row))
                if len(row) != len(names):
                    raise ConversionError(
                        f"A linha de dados {count} não tem a mesma quantidade de colunas do cabeçalho."
                    )
                values = []
                for name, value in zip(names, row, strict=True):
                    validate_cell(value, limits)
                    value = cast_value(value, types.get(name))
                    validate_cell(value, limits)
                    values.append(sql_literal(value, dialect))
                line = f"INSERT INTO {target} ({columns}) VALUES ({', '.join(values)});\n".encode(
                    "utf-8"
                )
                if output.tell() + len(line) > limits.output_bytes:
                    raise ConversionError(
                        "O SQL gerado excede o limite de saída. Divida o arquivo."
                    )
                output.write(line)
                if row_limit is not None and count == row_limit:
                    break
        budget.check()
    except ConversionError:
        raise
    except Exception:
        raise ConversionError(
            "Não foi possível ler o arquivo. Verifique o formato e a integridade dos dados."
        ) from None
    if count == 0:
        raise ConversionError("O arquivo não contém linhas de dados.")
    output.seek(0)
    return output, count
